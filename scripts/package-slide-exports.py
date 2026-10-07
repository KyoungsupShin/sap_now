"""Package the exact same slide renders into full-bleed PPTX and PDF."""
from pathlib import Path
import json
import sys
import fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.dml.color import RGBColor
import re
from urllib.parse import urlparse, unquote
from PIL import Image
from optimize_pptx import optimize_pptx

root = Path(__file__).resolve().parents[1]
editable_only = '--editable-only' in sys.argv
paths = [arg for arg in sys.argv[1:] if not arg.startswith('--')]
rendered = Path(paths[0]) if paths else root / 'exports' / 'rendered'
output = root / 'exports'
slides = json.loads((rendered / 'manifest.json').read_text())
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = 'SAP NOW — Autonomous Enterprise'
prs.core_properties.subject = 'Faithful export of the HTML presentation'
prs.core_properties.author = 'KyoungsupShin'
doc = fitz.open()
if not editable_only:
    for item in slides:
        png = rendered / f"{item['order']:02}.png"
        with Image.open(png) as image:
            assert image.size == (3200, 1800), (png, image.size)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = f"{item['order']:02} — {item['title']}\nHTML source: {item['file']}"
        page = doc.new_page(width=960, height=540)
        page.insert_image(page.rect, filename=str(png))
    prs.save(output / 'SAP_NOW_exact.pptx')
    optimize_pptx(output / 'SAP_NOW_exact.pptx')
# A second deck keeps original complex content but exposes flat card geometry.
hybrid = Presentation()
hybrid.slide_width, hybrid.slide_height = prs.slide_width, prs.slide_height
hybrid.core_properties.title = 'SAP NOW — Editable text and shapes'
def color(value):
    parts = [int(v) for v in re.findall(r'[\d.]+', value)[:3]]
    return RGBColor(*parts)
def unit(px):
    return Inches(px / 120)
def set_picture_background(slide, image_path):
    # Use DrawingML background fill rather than a selectable picture shape.
    _, relationship_id = slide.part.get_or_add_image_part(str(image_path))
    background = slide._element.cSld.get_or_add_bg()
    background_properties = background.get_or_add_bgPr()
    for child in list(background_properties):
        background_properties.remove(child)
    fill = OxmlElement('a:blipFill')
    blip = OxmlElement('a:blip')
    blip.set(qn('r:embed'), relationship_id)
    fill.append(blip)
    stretch = OxmlElement('a:stretch')
    stretch.append(OxmlElement('a:fillRect'))
    fill.append(stretch)
    background_properties.append(fill)
    background_properties.append(OxmlElement('a:effectLst'))
for item in slides:
    slide = hybrid.slides.add_slide(hybrid.slide_layouts[6])
    base = rendered / f"{item['order']:02}-base.png"
    set_picture_background(slide, base)
    for ribbon in item.get('nativeCurves', []):
        for curve in ribbon['curves']:
            builder=slide.shapes.build_freeform(0,0,scale=(unit(ribbon['width']/ribbon['viewWidth']),unit(ribbon['height']/ribbon['viewHeight'])))
            builder.add_line_segments([(ribbon['viewWidth'],0),(ribbon['viewWidth'],ribbon['viewHeight']),(0,ribbon['viewHeight'])],close=True)
            shape=builder.convert_to_shape(unit(ribbon['x']),unit(ribbon['y']))
            geom=shape._element.spPr.find(qn('a:custGeom'));paths=geom.find(qn('a:pathLst'))
            for child in list(paths):paths.remove(child)
            path=OxmlElement('a:path');path.set('w',str(round(ribbon['viewWidth']*100)));path.set('h',str(round(ribbon['viewHeight']*100)));path.set('fill','none');path.set('stroke','1')
            tokens=re.findall(r'[MCS]|-?\d+(?:\.\d+)?',curve['d']);i=0;current=(0,0);control=None
            def point(parent,x,y):
                pt=OxmlElement('a:pt');pt.set('x',str(round(x*100)));pt.set('y',str(round(y*100)));parent.append(pt)
            while i<len(tokens):
                command=tokens[i];i+=1
                if command=='M':
                    x,y=map(float,tokens[i:i+2]);i+=2;node=OxmlElement('a:moveTo');point(node,x,y);current=(x,y);control=None
                else:
                    count=6 if command=='C' else 4;v=list(map(float,tokens[i:i+count]));i+=count
                    if command=='S':v=[2*current[0]-control[0],2*current[1]-control[1]]+v
                    node=OxmlElement('a:cubicBezTo')
                    for n in range(0,6,2):point(node,v[n],v[n+1])
                    control=(v[2],v[3]);current=(v[4],v[5])
                path.append(node)
            paths.append(path);shape.fill.background();shape.line.color.rgb=RGBColor.from_string(curve['stroke'].lstrip('#'));shape.line.width=Pt(curve['width']*.6)
            shape.line._get_or_add_ln().set('cap','rnd')
            for style in list(shape._element.findall(qn('p:style'))):shape._element.remove(style)
    for card in item.get('nativeCards', []):
        x,y,w,h = (card[k] for k in ('x','y','width','height'))
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if card['radius'] else MSO_SHAPE.RECTANGLE
        rect = slide.shapes.add_shape(kind, unit(x), unit(y), unit(w), unit(h))
        for style in list(rect._element.findall(qn('p:style'))): rect._element.remove(style)
        if card['radius']:
            rect.adjustments[0] = min(0.5, card['radius'] / min(w,h))
        if card['fill'] in ('rgba(0, 0, 0, 0)', 'transparent'):
            rect.fill.background()
        else:
            rect.fill.solid(); rect.fill.fore_color.rgb = color(card['fill'])
        if card.get('edges') and len({(e['width'],e['color']) for e in card['edges']})>1:
            rect.line.fill.background()
            for edge in card['edges']:
                if not edge['width']: continue
                side=edge['side']; thickness=edge['width']
                ex=x if side!='Right' else x+w-thickness
                ey=y if side!='Bottom' else y+h-thickness
                ew=w if side in ('Top','Bottom') else thickness
                eh=thickness if side in ('Top','Bottom') else h
                border=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,unit(ex),unit(ey),unit(ew),unit(eh))
                border.fill.solid(); border.fill.fore_color.rgb=color(edge['color']); border.line.fill.background()
                for style in list(border._element.findall(qn('p:style'))):border._element.remove(style)
        elif card['borderWidth']:
            rect.line.color.rgb = color(card['border'])
            rect.line.width = Pt(card['borderWidth'] * .6)
        else:
            rect.line.fill.background()
        if card['topWidth'] and not card.get('edges'):
            top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, unit(x), unit(y), unit(w), unit(card['topWidth']))
            for style in list(top._element.findall(qn('p:style'))): top._element.remove(style)
            top.fill.solid(); top.fill.fore_color.rgb = color(card['topColor']); top.line.fill.background()
    for icon in item.get('nativeIcons', []):
        triangle = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, unit(icon['x']-5), unit(icon['y']+5), unit(icon['height']), unit(icon['width']))
        triangle.rotation = 90
        triangle.fill.solid(); triangle.fill.fore_color.rgb = RGBColor.from_string(icon['color'].lstrip('#'))
        triangle.line.fill.background()
        for style in list(triangle._element.findall(qn('p:style'))): triangle._element.remove(style)
    for media in item.get('nativeMedia', []):
        source=root / unquote(urlparse(media['src']).path.lstrip('/'))
        x,y,w,h=(media[k] for k in ('x','y','width','height'))
        if media['kind']=='video':
            slide.shapes.add_picture(media['preview'],unit(x),unit(y),unit(w),unit(h))
            movie=slide.shapes.add_movie(str(source),unit(x),unit(y),unit(w),unit(h),poster_frame_image=media['preview'],mime_type='video/mp4')
            if item['order']==2:
                for timing in slide._element.findall('.//' + qn('p:video')):
                    for node in timing.findall('.//' + qn('p:cTn')):
                        node.set('repeatCount','indefinite')
                    for cond in timing.findall('.//' + qn('p:cond')):
                        cond.set('delay','0')
        elif source.suffix.lower()=='.gif':
            frame=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,unit(x),unit(y),unit(w),unit(h))
            frame.adjustments[0]=.015;frame.fill.solid();frame.fill.fore_color.rgb=RGBColor(238,243,247);frame.line.color.rgb=RGBColor(197,215,232);frame.line.width=Pt(.6)
            nw,nh=media['naturalWidth'],media['naturalHeight'];ratio=min(w/nw,h/nh);fw,fh=nw*ratio,nh*ratio
            slide.shapes.add_picture(str(source),unit(x+(w-fw)/2),unit(y+(h-fh)/2),unit(fw),unit(fh))
        else:
            slide.shapes.add_picture(media['preview'],unit(x),unit(y),unit(w),unit(h))
    # Preserve each DOM text fragment's x position, including flex/inline gaps.
    groups = []
    for line in item.get('nativeText', []):
        group = next((g for g in groups if g['owner'] == line.get('owner')
                      and abs(g['y']-line['y']) < 1.5
                      and abs(max(r['x']+r['width'] for r in g['runs'])-line['x']) < 1.5), None)
        if group is None:
            group = {'owner':line.get('owner'), 'y':line['y'], 'runs':[]}
            groups.append(group)
        group['runs'].append(line)
    for group in groups:
        lines = sorted(group['runs'], key=lambda l:l['x'])
        left = min(l['x'] for l in lines)
        right = max(l['x']+l['width'] for l in lines)
        height = max(l['height'] for l in lines)
        text = slide.shapes.add_textbox(unit(left), unit(group['y']), unit(right-left+16), unit(height+8))
        frame = text.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        paragraph = frame.paragraphs[0]
        paragraph.space_before = paragraph.space_after = Pt(0)
        for line in lines:
            run = paragraph.add_run(); run.text = line['text']
            run.font.name = line['font']
            run._r.get_or_add_rPr().set('spc', str(round(line.get('letterSpacing', 0) * 60)))
            for tag in ('a:ea', 'a:cs'):
                font_element = OxmlElement(tag)
                font_element.set('typeface', 'Noto Sans CJK KR' if tag == 'a:ea' else line['font'])
                run._r.get_or_add_rPr().append(font_element)
            run.font.size = Pt(line['fontSize'] * .6)
            run.font.bold = int(line['weight']) >= 600
            run.font.italic = line['italic']; run.font.color.rgb = color(line['color'])
    slide.notes_slide.notes_text_frame.text = f"{item['order']:02} — {item['title']}\nAll HTML text is editable; card backgrounds and borders are native shapes. Text embedded in source images remains part of those images."
hybrid.save(output / 'SAP_NOW_editable.pptx')
optimize_pptx(output / 'SAP_NOW_editable.pptx')
if not editable_only:
    doc.set_metadata({'title': 'SAP NOW — Autonomous Enterprise', 'author': 'KyoungsupShin'})
    doc.save(output / 'SAP_NOW.pdf', garbage=4, deflate=True)
assert len(hybrid.slides) == len(slides)
print(f'Created editable PPTX: {len(slides)} pages, {sum(1 for slide in hybrid.slides for shape in slide.shapes if shape.has_text_frame and shape.text.strip())} text boxes.')
