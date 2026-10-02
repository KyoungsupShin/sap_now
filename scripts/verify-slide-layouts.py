"""Compare a LibreOffice-rendered editable PPTX PDF with the HTML captures.

Usage: python scripts/verify-slide-layouts.py /path/to/SAP_NOW_editable.pdf
Creates one paired review page per slide; pixel differences are diagnostic,
not a pass/fail claim about Microsoft PowerPoint rendering.
"""
from pathlib import Path
import io
import json
import sys
import hashlib
import fitz
import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
rendered = root / 'exports' / 'rendered'
out = rendered / 'review'
out.mkdir(exist_ok=True)
manifest = json.loads((rendered / 'manifest.json').read_text())
actual = fitz.open(sys.argv[1])
editable = Presentation(root / 'exports' / 'SAP_NOW_editable.pptx')
exact = Presentation(root / 'exports' / 'SAP_NOW_exact.pptx')
assert len(actual) == len(editable.slides) == len(exact.slides) == len(manifest)
review = fitz.open()
rows = []
for index, item in enumerate(manifest):
    number = item['order']
    source = rendered / f'{number:02}.png'
    assert hashlib.sha256(exact.slides[index].shapes[0].image.blob).digest() == hashlib.sha256(source.read_bytes()).digest()
    html = Image.open(source).convert('RGB').resize((1600, 900), Image.Resampling.LANCZOS)
    pix = actual[index].get_pixmap(matrix=fitz.Matrix(1600 / actual[index].rect.width, 1600 / actual[index].rect.width), alpha=False)
    ppt = Image.frombytes('RGB', (pix.width, pix.height), pix.samples).resize((1600, 900))
    ppt.save(out / f'{number:02}-pptx.png')
    delta = np.abs(np.asarray(html, dtype=np.int16) - np.asarray(ppt, dtype=np.int16))
    changed = float((delta.max(axis=2) > 32).mean() * 100)
    rows.append({'order': number, 'file': item['file'], 'changed_pixels_percent': round(changed, 3), 'mean_channel_difference': round(float(delta.mean()), 3)})
    pair = Image.new('RGB', (3200, 960), '#eaecee')
    draw = ImageDraw.Draw(pair)
    draw.text((20, 20), f'{number:02} HTML', fill='black')
    draw.text((1620, 20), f'{number:02} Editable PPTX / LibreOffice', fill='black')
    pair.paste(html, (0, 60)); pair.paste(ppt, (1600, 60))
    pair.save(out / f'{number:02}-comparison.jpg', quality=90)
    stream = io.BytesIO(); pair.save(stream, format='JPEG', quality=90)
    page = review.new_page(width=1600, height=480)
    page.insert_image(page.rect, stream=stream.getvalue())
for start in range(0, len(manifest), 5):
    sheet = Image.new('RGB', (1600, 2400), '#eaecee')
    for index in range(start, min(start + 5, len(manifest))):
        pair = Image.open(out / f'{index+1:02}-comparison.jpg')
        sheet.paste(pair.resize((1600, 480)), (0, (index-start)*480))
    sheet.save(out / f'contact-{start+1:02}.jpg', quality=92)
review.save(out / 'SAP_NOW_layout_review.pdf')
(out / 'metrics.json').write_text(json.dumps(rows, indent=2))
print(json.dumps(rows, indent=2))
print('All exact PPTX source images match HTML screenshots byte-for-byte.')
